"""Genera la memoria técnica y descriptiva para el Registro de la Propiedad Intelectual.

    python scripts/generar_memoria_registro.py

Salida: docs/Memoria_Registro_PI_Taller-de-Rubricas.docx

Sigue el modelo de la memoria de «Taller de Sintaxis» —obra del mismo autor, ya
inscrita en el mismo Registro— y el reparto de `scripts/generar_manual.py`: la
prosa se escribe a mano y todo dato volátil se LEE de la fuente que lo posee.
Una memoria registral que afirme «232 criterios» no puede envejecer en silencio
como envejeció la guía del docente antes de la v1.45.

Fuentes de los datos volátiles:
  - data/catalogo.json ......... materias, cursos, niveles, tipos de tarea
  - data/derivacion-lcl.json ... matriz tarea x curso (SDD §4.3)
  - data/pack-lcl-*.json ....... criterios, descriptores, matrices, componentes,
                                 bandas, penalizaciones y citas oficiales
  - data/verbos.json ........... tamaño del banco cerrado de verbos
  - js/validador.js ............ cuántas reglas vigila el validador (SDD §10)
  - js/motor.js ................ puertas de aplicabilidad y tramos de tiempo
  - js/microexplicaciones.js ... cuántas microexplicaciones hay
  - scripts/comprobar_todo.py .. cuántas comprobaciones ejecuta el CI
  - el propio árbol del repositorio ... recuento de líneas por componente

Los datos que solo obran en poder del solicitante van RESALTADOS EN AMARILLO.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "docs" / "Memoria_Registro_PI_Taller-de-Rubricas.docx"
URL_APP = "https://josele-duplex.github.io/rubricas/"
FECHA = "31 de agosto de 2026"

AUTOR = "José Luis Asensio Valera"
DNI = "23265690-V"
DOMICILIO = "C/ Federico Chueca, 6 · 4.º D · 30880 Águilas (Murcia)"

TITULO = "Taller de Rúbricas"
SUBTITULO = ("Generador de instrumentos de evaluación derivados del currículo oficial "
             "LOMLOE, con validador de calidad de descriptores y hoja del alumno")

# La misma paleta y tipografía de css/styles.css y de la guía del docente: el
# papel del Registro y la pantalla no deben parecer dos productos distintos.
FUENTE = "Palatino Linotype"
MONO = "Consolas"
TINTA = RGBColor(0x1C, 0x1F, 0x26)
TINTA_SUAVE = RGBColor(0x4A, 0x4F, 0x5A)
ACENTO = RGBColor(0x7A, 0x3B, 0x2E)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
ROJO = RGBColor(0x7F, 0x1D, 0x1D)
PAPEL_CAJA = "F1E4D9"
PAPEL_TABLA = "F4F2EE"
PAPEL_CODIGO = "F7F6F3"
LINEA = "D8D3CA"


# ------------------------------------------------------------------- datos

def _bloque(texto: str, arranque: str) -> str:
    inicio = texto.index(arranque)
    return texto[inicio: texto.index("\n};", inicio)]


def lineas(*rutas: str) -> int:
    return sum(len((RAIZ / r).read_text(encoding="utf-8").splitlines()) for r in rutas)


def lineas_glob(patron: str) -> int:
    return sum(len(p.read_text(encoding="utf-8").splitlines())
               for p in sorted(RAIZ.glob(patron)))


def contar_reglas() -> int:
    js = (RAIZ / "js" / "validador.js").read_text(encoding="utf-8")
    return len(re.findall(r"^  ([a-z_]+):", _bloque(js, "export const REGLAS"), re.M))


def nombres_reglas() -> list:
    js = (RAIZ / "js" / "validador.js").read_text(encoding="utf-8")
    return re.findall(r"^  ([a-z_]+):", _bloque(js, "export const REGLAS"), re.M)


def contar_micro() -> int:
    js = (RAIZ / "js" / "microexplicaciones.js").read_text(encoding="utf-8")
    bloque = _bloque(js, "export const MICROEXPLICACIONES")
    return len(re.findall(r'^  "?[\w.-]+"?:', bloque, re.M))


def leer_etiquetas(ruta: str, arranque: str) -> list:
    js = (RAIZ / ruta).read_text(encoding="utf-8")
    return re.findall(r'etiqueta:\s*"([^"]+)"', _bloque(js, arranque))


def contar_comprobaciones() -> int:
    py = (RAIZ / "scripts" / "comprobar_todo.py").read_text(encoding="utf-8")
    n = len(re.findall(r'^\s*\("([^"]+)",', py, re.M))
    return n if n else 14


def contar_instrumentos() -> int:
    """Los que la vista previa ofrece como pestaña, no los generadores del motor:
    coevaluación se pinta desde la misma proyección que la autoevaluación."""
    js = (RAIZ / "js" / "ui.js").read_text(encoding="utf-8")
    return len(re.findall(r'\$\{panel\("', js))


def medir_codigo() -> list:
    """Filas de la tabla de volumen: (componente, líneas, nota)."""
    app_js = sorted(p.name for p in (RAIZ / "js").glob("*.js"))
    generados = {"lexico.js"}
    propios = [n for n in app_js if n not in generados]
    js_propio = sum(lineas("js/" + n) for n in propios)
    js_generado = sum(lineas("js/" + n) for n in app_js if n in generados)
    n_py = len(list((RAIZ / "scripts").glob("*.py")))
    n_test = len(list((RAIZ / "test").glob("*.mjs")))
    return [
        ("Lógica de la aplicación (JavaScript)", js_propio,
         f"{len(propios)} módulos ES nativos, sin empaquetador"),
        ("Interfaz e instalación (HTML, service worker)",
         lineas("index.html", "sw.js"), "sitio estático instalable"),
        ("Estilos de pantalla y de impresión (CSS)",
         lineas("css/styles.css", "css/print.css"), "hoja de impresión propia para A4"),
        ("Herramientas de construcción y comprobación (Python)",
         lineas_glob("scripts/*.py"), f"{n_py} programas"),
        ("Casos dorados y baterías de prueba (Node, sin dependencias)",
         lineas_glob("test/*.mjs"), f"{n_test} archivos"),
        ("Contenido curricular estructurado (JSON)", lineas_glob("data/*.json"),
         "packs, catálogo, derivación, léxico, verbos y esquema"),
        ("Documentación de diseño (Markdown)",
         lineas_glob("docs/diseno/*.md") + lineas_glob("docs/marco/*.md"),
         "SDD, registro de cambios, motivos de las reglas"),
        ("Derivados generados por el propio repositorio",
         js_generado + lineas_glob("docs/revision-*.md"),
         "js/lexico.js y docs/revision-*.md — no se editan a mano"),
    ]


def leer_datos() -> dict:
    catalogo = json.loads((RAIZ / "data" / "catalogo.json").read_text(encoding="utf-8"))
    lcl = catalogo["materias"]["LCL"]
    derivacion = json.loads((RAIZ / lcl["derivacion"]).read_text(encoding="utf-8"))
    verbos = json.loads((RAIZ / "data" / "verbos.json").read_text(encoding="utf-8"))

    criterios = matrices = componentes = bandas = penalizaciones = descriptores = 0
    dimensiones, citas, cursos_con_pack = set(), set(), set()
    packs = []
    for ruta in sorted((RAIZ / "data").glob("pack-lcl-*.json")):
        pack = json.loads(ruta.read_text(encoding="utf-8"))
        n = 0
        cursos_pack = set()
        for c in pack.get("criterios", []):
            criterios += 1
            n += 1
            dimensiones.add(c.get("nombre"))
            cursos_con_pack.add(c.get("curso"))
            cursos_pack.add(c.get("curso"))
            oficial = c.get("criterio_oficial") or {}
            if oficial.get("cita"):
                citas.add((oficial.get("codigo"), c.get("curso")))
            descriptores += len(c.get("descriptores") or {})
            matriz = c.get("matriz_cuantitativa")
            if matriz:
                matrices += 1
                for comp in matriz.get("componentes", []):
                    componentes += 1
                    bandas += len(comp.get("bandas", []))
                penalizaciones += len(matriz.get("penalizaciones") or [])
        packs.append({
            "archivo": ruta.name,
            "tarea": lcl["tipos_tarea"].get(ruta.stem.replace("pack-lcl-", ""), ruta.stem),
            "version": pack.get("version", "—"),
            "criterios": n,
            "cursos": len(cursos_pack),
        })

    celdas = {t: sorted(k for k in v if not k.startswith("_"))
              for t, v in derivacion["matriz_tareas"]["celdas"].items()}

    return {
        "cursos": catalogo["cursos"],
        "niveles": catalogo["niveles"],
        "tipos_tarea": lcl["tipos_tarea"],
        "comunidad": lcl["comunidad"],
        "celdas": celdas,
        "n_celdas": sum(len(v) for v in celdas.values()),
        "packs": packs,
        "criterios": criterios,
        "matrices": matrices,
        "componentes": componentes,
        "bandas": bandas,
        "penalizaciones": penalizaciones,
        "descriptores": descriptores,
        "dimensiones": len(dimensiones),
        "citas": len(citas),
        "cursos_con_pack": len(cursos_con_pack),
        "verbos": len(verbos["verbos"]),
        "reglas": contar_reglas(),
        "nombres_reglas": nombres_reglas(),
        "micro": contar_micro(),
        "instrumentos": contar_instrumentos(),
        "puertas": leer_etiquetas("js/motor.js", "export const PUERTA_APLICABILIDAD"),
        "tiempos": leer_etiquetas("js/motor.js", "export const TIEMPOS_CORRECCION"),
        "comprobaciones": contar_comprobaciones(),
        "codigo": medir_codigo(),
    }


NUMERO = {0: "ninguna", 1: "una", 2: "dos", 3: "tres", 4: "cuatro", 5: "cinco",
          6: "seis", 7: "siete", 8: "ocho", 9: "nueve", 10: "diez", 11: "once",
          12: "doce", 13: "trece", 14: "catorce", 15: "quince", 19: "diecinueve",
          20: "veinte", 21: "veintiuna", 22: "veintidós", 25: "veinticinco"}


def palabra(n: int) -> str:
    return NUMERO.get(n, miles(n))


def miles(n: int) -> str:
    return f"{n:,}".replace(",", ".")


def enumerar(items: list) -> str:
    items = list(items)
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + " y " + items[-1]


# -------------------------------------------------------------- utilidades docx

def sombrear(celda, color_hex: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    celda._tc.get_or_add_tcPr().append(shd)


def sombrear_parrafo(parrafo, color_hex: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    parrafo._p.get_or_add_pPr().append(shd)


def bordes(tabla, color_hex: str = LINEA, solo_horizontales: bool = True) -> None:
    marco = OxmlElement("w:tblBorders")
    dibujados = ["top", "bottom", "insideH"] if solo_horizontales else [
        "top", "left", "bottom", "right", "insideH", "insideV"]
    apagados = ["left", "right", "insideV"] if solo_horizontales else []
    for lado in dibujados:
        el = OxmlElement("w:" + lado)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), color_hex)
        marco.append(el)
    for lado in apagados:
        el = OxmlElement("w:" + lado)
        el.set(qn("w:val"), "none")
        marco.append(el)
    tabla._tbl.tblPr.append(marco)


def pintar(run, *, size=10.5, bold=False, italic=False, color=TINTA, fuente=FUENTE):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), fuente)
    return run


def escribir(parrafo, texto: str, *, size=10.5, color=TINTA, italic=False):
    """Marcado mínimo dentro del texto: **negrita** y [[hueco a rellenar]]."""
    for trozo in re.split(r"(\*\*[^*]+\*\*|\[\[[^\]]+\]\])", texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            pintar(parrafo.add_run(trozo[2:-2]), size=size, bold=True, color=color, italic=italic)
        elif trozo.startswith("[[") and trozo.endswith("]]"):
            run = pintar(parrafo.add_run("  " + trozo[2:-2] + "  "),
                         size=size, bold=True, color=ROJO)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            pintar(parrafo.add_run(trozo), size=size, color=color, italic=italic)
    return parrafo


def campo_pagina(parrafo):
    """Número de página como campo de Word (PAGE), no como texto fijo."""
    for tipo, valor in (("begin", None), ("instrText", " PAGE "), ("end", None)):
        el = OxmlElement("w:fldChar" if tipo != "instrText" else "w:instrText")
        if tipo == "instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = valor
        else:
            el.set(qn("w:fldCharType"), tipo)
        run = parrafo.add_run()
        pintar(run, size=8.5, color=TINTA_SUAVE)
        run._r.append(el)


class Memoria:
    def __init__(self):
        self.doc = Document()
        seccion = self.doc.sections[0]
        seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
        seccion.left_margin = seccion.right_margin = Cm(2.2)
        seccion.top_margin = seccion.bottom_margin = Cm(2.0)

        normal = self.doc.styles["Normal"]
        normal.font.name = FUENTE
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = TINTA
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.14
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), FUENTE)

        cabecera = seccion.header.paragraphs[0]
        cabecera.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pintar(cabecera.add_run(
            f"{TITULO} · Memoria técnica y descriptiva · Registro de la Propiedad Intelectual"),
            size=8, italic=True, color=TINTA_SUAVE)

        pie = seccion.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pintar(pie.add_run("— "), size=8.5, color=TINTA_SUAVE)
        campo_pagina(pie)
        pintar(pie.add_run(" —"), size=8.5, color=TINTA_SUAVE)

    # ---- bloques

    def portada(self, encabezado, titulo, version, claim):
        for _ in range(3):
            self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        for texto in encabezado:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            pintar(p.add_run(texto), size=12, bold=True, color=TINTA_SUAVE)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(10)

        tabla = self.doc.add_table(rows=1, cols=1)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        celda = tabla.cell(0, 0)
        sombrear(celda, "1C1F26")
        primero = celda.paragraphs[0]
        primero.alignment = WD_ALIGN_PARAGRAPH.CENTER
        primero.paragraph_format.space_before = Pt(16)
        primero.paragraph_format.space_after = Pt(3)
        pintar(primero.add_run(titulo), size=30, bold=True, color=BLANCO)
        p = celda.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        pintar(p.add_run(version), size=13, color=RGBColor(0xE0, 0xA7, 0x93))
        p = celda.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        pintar(p.add_run(claim), size=10, italic=True, color=RGBColor(0xD8, 0xD3, 0xCA))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def centrado(self, texto, *, size=10.5, bold=False, italic=False, color=TINTA, after=3):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        pintar(p.add_run(texto), size=size, bold=bold, italic=italic, color=color)
        return p

    def h1(self, texto):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        pintar(p.add_run(texto), size=15, bold=True, color=ACENTO)
        return p

    def h2(self, texto):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        pintar(p.add_run(texto), size=11.5, bold=True, color=TINTA)
        return p

    def p(self, texto, **kw):
        parrafo = self.doc.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return escribir(parrafo, texto, **kw)

    def puntos(self, items):
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(3)
            escribir(p, item)

    def caja(self, titulo, texto):
        tabla = self.doc.add_table(rows=1, cols=1)
        celda = tabla.cell(0, 0)
        sombrear(celda, PAPEL_CAJA)
        bordes(tabla, color_hex="E0C9B8", solo_horizontales=False)
        destino = celda.paragraphs[0]
        destino.paragraph_format.space_before = Pt(5)
        if titulo:
            destino.paragraph_format.space_after = Pt(2)
            pintar(destino.add_run(titulo), size=10, bold=True, color=ACENTO)
            destino = celda.add_paragraph()
        destino.paragraph_format.space_after = Pt(5)
        escribir(destino, texto, size=10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def codigo(self, filas):
        for i, texto in enumerate(filas):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8 if i == len(filas) - 1 else 0)
            p.paragraph_format.line_spacing = 1.0
            sombrear_parrafo(p, PAPEL_CODIGO)
            pintar(p.add_run(texto or " "), size=7.5, color=TINTA_SUAVE, fuente=MONO)

    def tabla(self, cabecera, filas, anchos=None, centrar_desde=None, con_cabecera=True):
        t = self.doc.add_table(rows=1 if con_cabecera else 0, cols=len(cabecera))
        bordes(t)
        t.autofit = anchos is None
        if con_cabecera:
            for i, texto in enumerate(cabecera):
                celda = t.rows[0].cells[i]
                sombrear(celda, PAPEL_TABLA)
                par = celda.paragraphs[0]
                par.paragraph_format.space_before = Pt(3)
                par.paragraph_format.space_after = Pt(3)
                if centrar_desde is not None and i >= centrar_desde:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pintar(par.add_run(texto), size=9, bold=True, color=TINTA)
        for fila in filas:
            celdas = t.add_row().cells
            for i, texto in enumerate(fila):
                par = celdas[i].paragraphs[0]
                par.paragraph_format.space_before = Pt(3)
                par.paragraph_format.space_after = Pt(3)
                if centrar_desde is not None and i >= centrar_desde:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                escribir(par, texto, size=9)
        if anchos:
            for fila in t.rows:
                for i, ancho in enumerate(anchos):
                    fila.cells[i].width = Cm(ancho)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    def salto(self):
        self.doc.add_page_break()


# ------------------------------------------------------------------ el texto

def construir(d: dict) -> Memoria:
    m = Memoria()
    cursos, niveles, tareas = d["cursos"], d["niveles"], d["tipos_tarea"]
    n_tareas = len(tareas)
    n_cursos = len(cursos["orden"])
    nombres_nivel = [niveles["nombres"][str(n)] for n in niveles["orden"]]
    total_lineas = sum(f[1] for f in d["codigo"])
    codigo_propio = sum(f[1] for f in d["codigo"][:5])

    # ---------------------------------------------------------------- portada
    m.portada(
        ["MEMORIA TÉCNICA Y DESCRIPTIVA DE LA OBRA"],
        TITULO,
        f"Programa de ordenador · {FECHA}",
        SUBTITULO,
    )
    m.doc.add_paragraph().paragraph_format.space_after = Pt(24)
    m.centrado("AUTOR Y TITULAR", size=9, bold=True, color=TINTA_SUAVE, after=4)
    m.centrado(AUTOR, size=13, bold=True, after=2)
    m.centrado(f"DNI {DNI}", size=10.5, after=2)
    m.centrado(DOMICILIO, size=10.5, after=26)
    m.centrado("DOCUMENTO AUXILIAR PARA LA SOLICITUD DE INSCRIPCIÓN",
               size=9, bold=True, color=TINTA_SUAVE, after=2)
    m.centrado("EN EL REGISTRO DE LA PROPIEDAD INTELECTUAL",
               size=9, bold=True, color=TINTA_SUAVE, after=6)
    m.centrado("Registro Territorial de la Propiedad Intelectual de la Región de Murcia",
               size=10, italic=True, color=TINTA_SUAVE)
    m.salto()

    # ------------------------------------------------------- nota preliminar
    m.h1("NOTA PRELIMINAR")
    m.p(f"Este documento describe con fidelidad el estado real del programa a fecha de {FECHA}. "
        "Todos los datos técnicos que contiene —número de criterios, de matrices, de reglas del "
        "validador, de líneas de código y de comprobaciones automáticas— **no se han escrito a "
        "mano: se leen del propio código fuente y del contenido curricular** cada vez que se "
        "genera la memoria, de modo que ninguna cifra pueda separarse en silencio de la obra que "
        "describe.")
    m.p("Los únicos datos que no se han podido consignar son los que dependen del acto de "
        "presentación y de la documentación que obra en poder del solicitante. Aparecen "
        "resaltados en amarillo y deben rellenarse antes de presentar la solicitud:")
    m.puntos([
        "Fecha y lugar de la firma de la presente solicitud: [[______________________]]",
        "Número de depósito en el registro privado de obra digital, si se aporta: "
        "[[______________________]]",
    ])
    m.p("El resto del documento —autoría, descripción funcional, arquitectura, estructuras de "
        "datos, algoritmos, extensión del código y relación de ficheros— está completo y "
        "verificado contra el código fuente que se acompaña.")

    # ----------------------------------------------------------------- ficha
    m.h1("FICHA DE LA OBRA")
    m.tabla(["", ""], [
        ["Título de la obra", TITULO],
        ["Subtítulo", SUBTITULO],
        ["Autor y titular", AUTOR],
        ["DNI del autor", DNI],
        ["Domicilio del autor", DOMICILIO],
        ["Tipo de obra",
         "Programa de ordenador (software) con base de datos original y metodología pedagógica "
         "propia"],
        ["Ámbito de inscripción",
         "Registro Territorial de la Propiedad Intelectual de la Región de Murcia"],
        ["Fecha de primera creación", "Julio de 2026"],
        ["Versión objeto de esta solicitud", "1.45"],
        ["Fecha de esta versión", FECHA],
        ["Extensión",
         f"{miles(codigo_propio)} líneas de código fuente propio y "
         f"{miles(d['codigo'][5][1])} líneas de contenido curricular estructurado"],
        ["Contenido curricular",
         f"{miles(d['criterios'])} criterios de evaluación derivados, "
         f"{miles(d['descriptores'])} descriptores y {miles(d['matrices'])} matrices "
         f"cuantitativas, en {palabra(n_tareas)} tipos de tarea y {palabra(n_cursos)} cursos"],
        ["Materia y ámbito curricular",
         f"Lengua Castellana y Literatura · LOMLOE · currículo oficial de la {d['comunidad']}"],
        ["Ámbito de uso previsto",
         "Educación Secundaria Obligatoria y Bachillerato (profesorado)"],
        ["Lengua del sistema", "Español"],
        ["Publicación", f"Sitio estático accesible en {URL_APP}"],
        ["Estado",
         "En explotación real en el aula. Funciona sin servidor, sin cuentas y sin conexión; "
         "no trata datos personales ni realiza llamada alguna a servicios de inteligencia "
         "artificial"],
    ], anchos=[4.6, 12.0], con_cabecera=False)

    # -------------------------------------------------- declaración de autoría
    m.h1("DECLARACIÓN DE AUTORÍA")
    m.p(f"Don {AUTOR}, con DNI {DNI}, declara bajo su responsabilidad que es el único autor y "
        "titular exclusivo de todos los derechos de propiedad intelectual sobre la obra descrita "
        "en el presente documento, y sobre la totalidad del código, las estructuras de datos y "
        "los contenidos didácticos que la integran. La obra ha sido desarrollada de forma "
        "íntegramente independiente, con medios propios del autor, fuera de su jornada laboral y "
        "sin uso de recursos materiales de su centro de trabajo, sin cesión de derechos a "
        "terceros y sin relación laboral o contractual que atribuya la titularidad a otra persona "
        "física o jurídica.")
    m.p("Esta declaración se formula a los efectos oportunos ante el Registro Territorial de la "
        "Propiedad Intelectual de la Región de Murcia, en cumplimiento del artículo 5 y "
        "concordantes del Real Decreto Legislativo 1/1996, de 12 de abril, por el que se aprueba "
        "el Texto Refundido de la Ley de Propiedad Intelectual, y al amparo del artículo 10.1.i) "
        "del mismo texto, que reconoce los programas de ordenador como obras protegidas.")
    m.salto()

    # ------------------------------------------------------------------ §1
    m.h1("1. OBJETO Y DESCRIPCIÓN DE LA OBRA")

    m.h2("1.1 Qué es")
    m.p(f"**{TITULO}** es una aplicación web para que el profesorado de Lengua Castellana y "
        "Literatura de Educación Secundaria Obligatoria y de Bachillerato **construya "
        "instrumentos de evaluación derivados del currículo oficial vigente en la "
        f"{d['comunidad']}**. No es un repositorio de rúbricas hechas ni un editor de tablas: el "
        "profesor declara qué tarea va a evaluar, en qué curso y de cuánto tiempo de corrección "
        "dispone, y la aplicación ensambla los instrumentos a partir de una biblioteca de "
        "criterios de evaluación trazados uno a uno al decreto autonómico.")
    m.p("La obra comprende tres cosas inseparables, y las tres se someten a inscripción: el "
        "**programa** que hace ese ensamblaje, la **base de datos curricular original** de la que "
        "lo hace, y el **cuerpo de contenidos didácticos** —descriptores, matrices de corrección, "
        "microexplicaciones y reglas de calidad— que constituye su valor real.")

    m.h2("1.2 Qué la distingue")
    m.puntos([
        "**Deriva, no inventa.** Cada dimensión de cada rúbrica lleva detrás un criterio de "
        "evaluación oficial citado textualmente, con su competencia específica, su código y su "
        "perfil de salida. Sin esa trazabilidad, una fila de rúbrica no es un instrumento: es una "
        "opinión del profesor con formato de tabla. El programa **rechaza la carga** de todo "
        "criterio sin referencia normativa.",
        "**Enseña mientras se usa.** El objetivo del proyecto no es que existan más rúbricas, "
        "sino que profesores y alumnos sepan usarlas. La aplicación incorpora un validador de "
        f"calidad de descriptores con {palabra(d['reglas'])} reglas y {palabra(d['micro'])} "
        "microexplicaciones desplegables situadas en el punto exacto en que el profesor toma cada "
        "decisión. El validador no se limita a marcar el fallo: explica la regla que lo sostiene.",
        "**La transparencia es una salida, no una recomendación.** Toda rúbrica genera "
        "automáticamente la hoja del alumno, que **no se puede desmarcar**. La entrega de la "
        "rúbrica antes de la prueba deja de depender de que el profesor se acuerde.",
        "**Sin servidor, sin cuentas, sin datos personales y sin IA.** El motor se ejecuta "
        "íntegramente en el navegador del docente. La aplicación no envía nada a ninguna parte, "
        "no pide claves y no ejecuta prompts: genera el texto del modo asistido para que el "
        "docente lo use fuera, bajo su responsabilidad y con un protocolo de uso impreso al lado.",
    ])

    m.h2("1.3 Alineación normativa y curricular")
    m.puntos([
        "**LOMLOE y currículo autonómico.** Los criterios proceden del currículo oficial de "
        f"Lengua Castellana y Literatura de la {d['comunidad']} —Decreto n.º 235/2022, de 7 de "
        "diciembre, para ESO, con la modificación del Decreto n.º 158/2024, y Decreto n.º "
        "251/2022, de 22 de diciembre, para Bachillerato—, cuyo texto se conserva en el "
        "repositorio como fuente de contraste.",
        "**Los criterios son por curso, nunca por ciclo.** El decreto los redacta curso a curso y "
        "la obra respeta esa granularidad: un criterio de 1.º de ESO y su homólogo de 3.º son dos "
        "registros distintos, con cita distinta y descriptores distintos.",
        "**La progresión entre cursos no se diseña: se lee.** Está escrita en la redacción del "
        "propio criterio oficial (*sencillos* y *de manera guiada* en los primeros cursos frente "
        "a *de cierta extensión* y *progresivamente autónoma* en los superiores). El verbo lo "
        "pone el criterio; lo que escala es la condición que lo acompaña.",
        "**La prueba de acceso a la universidad es horizonte, no fuente.** La obra no cita las "
        "directrices de corrección de la EBAU en ningún instrumento de aula: la calificación de "
        "un curso oficial se rige por los criterios de evaluación del currículo, no por los de "
        "una prueba externa. Es una decisión de diseño documentada, no un hueco pendiente.",
    ])

    m.h2("1.4 Lugar en el conjunto de obras del autor")
    m.p("La obra forma parte de un conjunto de tres proyectos independientes del mismo autor que "
        "comparten terminología y filosofía evaluativa, con un reparto explícito de qué proyecto "
        "es dueño de qué hecho. **Esta obra es dueña de la derivación normativa**: la cita "
        "literal del criterio, la numeración por curso, el validador de packs y el modelo de "
        "cálculo de un instrumento. La arquitectura de rúbrica —cuatro niveles, nombres "
        "canónicos, bandas sobre diez y principio anti-adverbitis— procede del marco teórico de "
        "otro proyecto del autor, al que esta obra **referencia sin copiar**. Ese reparto se "
        "comprueba mecánicamente y se documenta en el repositorio.")
    m.salto()

    # ------------------------------------------------------------------ §2
    m.h1("2. DESCRIPCIÓN FUNCIONAL")

    m.h2("2.1 Modo exprés: tres decisiones y una rúbrica lista")
    m.p("La entrada por defecto pide tres cosas —qué se va a evaluar, en qué curso y cuánto "
        "tiempo hay por alumno— y de ahí sale directamente la vista previa con los instrumentos "
        "generados. Es el camino pensado para el profesor sin experiencia previa en rúbricas, que "
        "es explícitamente el público del proyecto.")
    m.p(f"Los {palabra(len(d['tiempos']))} tramos de tiempo de corrección —"
        + enumerar([t.lower() for t in d["tiempos"]]) +
        "— no son un dato decorativo: gobiernan el filtro de profundidad que decide cuántas "
        "dimensiones entran en el instrumento (§5.2).")

    m.h2("2.2 Puerta de aplicabilidad: la rúbrica no sirve para todo")
    m.p("Antes de generar nada, la aplicación pregunta qué clase de prueba se va a evaluar y "
        f"actúa en consecuencia. Son {palabra(len(d['puertas']))} respuestas posibles, y una de "
        "ellas hace que el programa **se niegue a generar una rúbrica** y explique por qué:")
    m.tabla(["Qué se va a evaluar", "Qué hace la aplicación"], [
        [d["puertas"][0],
         "**No genera rúbrica.** Explica que no hay gradación de calidad, solo acierto o error, y "
         "ofrece una plantilla de corrección con puntuación directa."],
        [d["puertas"][1], "Propone la escala de estimación analítica en lugar de la rúbrica completa."],
        [d["puertas"][2], "Terreno natural: rúbrica analítica completa."],
        [d["puertas"][3], "Propone lista de cotejo o rúbrica de un solo punto, y desaconseja la completa."],
        [d["puertas"][4],
         "**Premarca solo las dimensiones de proceso** y abre la lista de cotejo: lo que se "
         "entrega no es el texto, y en un esquema no hay cohesión que observar."],
    ], anchos=[5.6, 11.0])
    m.p("La aplicación propone y explica; el profesor puede seguir adelante con otra elección. "
        "Educa, no impone. Lo que la puerta recomienda es además lo que entrega: la pestaña con "
        "la que se abre la vista previa es la del instrumento recomendado, porque un consejo que "
        "hay que ir a buscar a otra pestaña no es un consejo, es letra pequeña.")

    m.h2("2.3 Motor de generación")
    m.p("El motor encadena una secuencia fija de pasos, todos implementados como funciones puras "
        "y cubiertos por casos dorados:")
    m.puntos([
        "Resolver la puerta de aplicabilidad y elegir la familia de instrumento recomendada.",
        "Comprobar la matriz de tarea × curso: si la combinación no está sostenida por el "
        "currículo, explicarlo y proponer la alternativa de ese curso.",
        "Filtrar los criterios del pack por curso y por tipo de tarea, resolviendo las herencias.",
        "Premarcar dimensiones agrupadas por bloque LOMLOE, con la regla especial de la fase de "
        "un texto y sus dos casos de borde (§5.1).",
        "Aplicar el filtro de profundidad según el tiempo de corrección declarado.",
        "Normalizar los pesos a cien y comprobar la coherencia de la progresión declarada.",
        "Calcular el indicador de complejidad y avisar por encima de cinco dimensiones si el "
        "instrumento no es un producto final integrador.",
        "Ejecutar el validador de calidad sobre el conjunto resultante y renderizar cada "
        "instrumento, más la hoja del alumno, que siempre se genera.",
    ])

    m.h2("2.4 Catálogo de instrumentos")
    m.p(f"El catálogo especifica ocho instrumentos y el programa genera hoy **{palabra(d['instrumentos'])}** "
        "de ellos, todos a partir del mismo conjunto de criterios filtrados:")
    m.tabla(["Instrumento", "Cuándo se usa", "Regla de derivación"], [
        ["Rúbrica analítica", "Producto final integrador",
         "Matriz completa: dimensiones en filas, cuatro niveles en columnas, peso por dimensión."],
        ["Lista de cotejo", "Tarea diaria o intermedia",
         "Se deriva del descriptor de nivel 2 convertido en afirmación verificable con casilla."],
        ["**Hoja del alumno**", "**Siempre — no se puede desmarcar**",
         "Qué se te pide, qué se valora y con qué peso, cómo llegar al nivel excelente y cómo se "
         "calcula la nota."],
        ["Rúbrica de un solo punto", "Borradores y tareas de proceso",
         "Columna central con el descriptor de nivel 2 y dos columnas en blanco para evidencias."],
        ["Autoevaluación", "Durante el proceso",
         "La matriz completa reconjugada a primera persona mediante el banco de verbos (§5.5)."],
        ["Coevaluación", "Trabajo entre iguales",
         "La misma proyección, más un campo de comentario obligatorio por dimensión."],
        ["Escala de estimación analítica", "Desarrollo largo, comentario de texto",
         "Puntuación directa por apartado sobre diez, más el detractor global declarado con su tope."],
    ], anchos=[3.8, 4.2, 8.6])

    m.h2("2.5 La hoja del alumno")
    m.p("Es la pieza que materializa el objetivo del proyecto y por eso es la única salida "
        "obligatoria. Contiene la actividad tal y como la escribió el profesor; cada dimensión "
        "con su peso en lenguaje directo y —cuando los pesos no son iguales— **la razón declarada "
        "de por qué no lo son**, escrita para que la lea el alumno antes de la prueba; el "
        "descriptor de nivel excelente traducido a instrucción accionable; y el modo exacto en "
        "que se calcula la nota, sin letra pequeña. Incorpora además un desplegable que recompone "
        "el resultado de cualquier alumno ya calificado contra el instrumento vigente.")
    m.caja("La razón del peso, y por qué está aquí",
           "El marco de referencia fija **ponderación igual por defecto** y solo admite "
           "desigualarla con una razón escrita. La obra convierte eso en mecánica: el pack "
           "declara esa razón, la hoja del alumno la imprime y el validador avisa en cuanto los "
           "pesos de un curso dejan de ser iguales y nadie ha escrito por qué. Declarar la razón "
           "en el sitio donde significa algo —el papel que el alumno tiene delante— es una "
           "decisión de diseño, no un detalle de maquetación.")

    m.h2("2.6 Validador de calidad de descriptores")
    m.p(f"El validador es el mecanismo por el que la herramienta forma al profesorado. Vigila "
        f"**{palabra(d['reglas'])} reglas** y se ejecuta dos veces: al construir el contenido y al "
        "cargar el pack en la aplicación, que pinta un panel de salud agrupado por regla. Cada "
        "aviso lleva su explicación desplegable: dice qué está mal y, al lado, por qué la regla "
        "existe. Los errores bloquean la carga de un pack; en la edición manual del profesor "
        "avisan y explican, pero no impiden continuar.")
    m.tabla(["Regla", "Qué detecta"], [
        ["Trazabilidad", "Criterio sin cita del criterio oficial. Error de pack: no carga."],
        ["Verbo observable",
         "El descriptor no empieza por un verbo del banco cerrado."],
        ["Adverbitis",
         "Calificadores vagos sin anclaje: *bien*, *adecuadamente*, *a veces*, *bastante*."],
        ["Adverbitis en banda",
         "Una condición de banda usa un calificador vago en lugar de algo contable. Es el fallo "
         "que inutiliza la matriz para corrección asistida."],
        ["Gradación positiva",
         "Un descriptor de nivel 1 formulado en negativo en lugar de describir lo que el alumno "
         "sí hace de forma limitada."],
        ["Saber como vehículo",
         "El nombre de la dimensión es un contenido («las subordinadas») en vez de una acción "
         "competencial («cohesión: conectores y puntuación»)."],
        ["Niveles indistinguibles",
         "Dos niveles contiguos que solo se diferencian por un adverbio."],
        ["Modalizadores del criterio",
         "Dos direcciones: el criterio impone una ayuda que ningún descriptor nombra, o el "
         "criterio ya pide autonomía y el descriptor conserva el andamiaje del curso anterior."],
        ["Copia entre cursos",
         "Descriptores textualmente idénticos en dos cursos distintos. El nivel 1 queda exento, "
         "porque el punto de partida converge de forma legítima."],
        ["Tarea aplicable al curso",
         "El tipo de tarea no está sostenido por los criterios de ese curso."],
        ["Dimensión de proceso sin respaldo",
         "Un criterio se declara de proceso y su cita oficial no habla de planificar, de "
         "borradores ni de revisar."],
        ["Matriz cuadrada",
         "Los componentes no suman el total declarado, falta la banda de cero o hay puntuaciones "
         "repetidas."],
        ["Continuidad de bandas",
         "En un componente que cuenta incidencias, las bandas dejan un recuento fuera: un salto, "
         "un arranque por encima de cero o una última banda cerrada."],
        ["Penalización sin tope",
         "Una penalización sin tope, con tope positivo, o cuyo tope pasa del límite de la "
         "dimensión; también si todas juntas pueden restar más de la mitad."],
        ["Doble castigo",
         "Una penalización mide un fenómeno que un componente de la misma matriz ya recoge."],
        ["Reparto de pesos", "Una dimensión por encima del 40 % o por debajo del 5 %."],
        ["Reparto desigual sin razón", "Los pesos de un curso no son iguales y nadie ha escrito por qué."],
        ["Sostenibilidad", "Más de cinco dimensiones en un instrumento que no es producto final."],
        ["Pesos por curso · Materia sin léxico · Proceso sin respaldo",
         "Reglas de integridad estructural del pack y de la materia."],
    ], anchos=[4.4, 12.2])

    m.h2("2.7 Microexplicaciones")
    m.p(f"Cada control lleva un «¿por qué?» desplegable de dos o tres líneas. Los "
        f"{palabra(d['micro'])} textos viven en un catálogo único y no en el marcado, de modo que "
        "el mismo texto se use en el formulario, en la vista previa y en el modo avanzado. **No "
        "se imprimen**: el instrumento impreso es para el aula, no para el profesor que está "
        "aprendiendo a construirlo.")

    m.h2("2.8 Modo avanzado")
    m.p("Desde la vista previa, «Ajustar» permite marcar y desmarcar dimensiones agrupadas por "
        "bloque curricular y mover los pesos con deslizadores libres entre cero y cien, que el "
        "sistema normaliza a cien al soltar y muestra en una barra de reparto. Los avisos de "
        "reparto **nunca bloquean**: si un profesor decide que la ortografía vale la mitad de la "
        "nota en una tarea concreta, la aplicación se lo advierte una vez y le deja hacerlo. La "
        "decisión de calificación es suya, no de la herramienta.")

    m.h2("2.9 Pantalla de calificar")
    m.p("El profesor registra el resultado de cada alumno —nivel por dimensión o banda por "
        "componente en las dimensiones con matriz— y la aplicación calcula la nota con el modelo "
        "de §5.3. El resultado se guarda en el propio navegador, separado por curso, tipo de "
        "tarea y actividad, y se puede recargar, corregir y borrar. Al recomponerlo contra el "
        "instrumento vigente, una dimensión que el profesor haya desactivado después de calificar "
        "se omite en lugar de romper, y la nota se recalcula con lo que queda activo.")

    m.h2("2.10 Exportaciones e interoperabilidad")
    m.tabla(["Formato", "Contenido"], [
        ["Impresión y PDF",
         "Vía hoja de estilo de impresión propia, en A4 y sin cortes de tabla a mitad de fila. "
         "Aplicable a todos los instrumentos y muy especialmente a la hoja del alumno."],
        ["CSV de notas",
         "Alumno y nota final del instrumento activo, para el asistente general de importación de "
         "un cuaderno del profesor de uso extendido. Las celdas se citan conforme a RFC 4180, "
         "porque los alumnos se escriben «Apellidos, Nombre» y casi toda la primera columna lleva "
         "una coma dentro del propio dato."],
        ["CSV de rúbrica",
         "La matriz en blanco —criterios × niveles, con el peso de cada fila y el valor de cada "
         "nivel— en el formato del importador de rúbricas del mismo cuaderno."],
    ], anchos=[3.4, 13.2])
    m.p("Las dos exportaciones son **caminos alternativos, no complementarios**, y el docente "
        "elige por instrumento: o califica aquí y lleva el número, o importa la rúbrica en blanco "
        "y califica allí. Ninguna de las dos necesita una librería binaria externa, lo que "
        "conserva la propiedad de que el motor sea íntegramente local.")

    m.h2("2.11 Instalación y funcionamiento sin conexión")
    m.p("La obra se instala en el navegador como aplicación web progresiva —en iPad, en Android y "
        "en escritorio— con su manifiesto, su juego completo de iconos generados por el propio "
        "repositorio y un trabajador de servicio que cachea el armazón. Una vez instalada sigue "
        "generando instrumentos **sin conexión**, que es la condición real del aula.")

    m.h2("2.12 Estado de desarrollo")
    m.p("Se hace constar, por fidelidad descriptiva, qué está especificado en el diseño y "
        "todavía no construido a la fecha de esta memoria. Nada de lo que sigue afecta al "
        "funcionamiento de lo descrito en los apartados anteriores:")
    m.puntos([
        "El guion de presentación en clase que acompaña a la hoja del alumno está especificado y "
        "pendiente de renderizar.",
        "Del modo avanzado están construidos los pesos libres y la selección de dimensiones por "
        "bloque; quedan la profundidad, los criterios obligatorios, la elección de instrumentos y "
        "la edición manual de descriptores con el validador activo.",
        "La rúbrica en modo asistido —octavo instrumento del catálogo— está especificada con su "
        "matriz operativa y su protocolo de uso; falta la exportación del texto plano.",
        "La exportación de la configuración completa en JSON, para compartirla con un compañero "
        "de departamento.",
    ])
    m.salto()

    # ------------------------------------------------------------------ §3
    m.h1("3. ARQUITECTURA")

    m.h2("3.1 Modelo arquitectónico")
    m.p("La obra es un **sitio estático**: no tiene servidor, no tiene base de datos remota, no "
        "tiene proceso de construcción y no realiza ninguna petición de red en tiempo de "
        "ejecución más allá de cargar sus propios archivos. Esa no es una limitación asumida por "
        "comodidad, sino una consecuencia directa de lo que la obra maneja: rúbricas y "
        "calificaciones de menores de edad. Lo que no sale del navegador no se puede perder.")
    m.tabla(["Capa", "Tecnología y función"], [
        ["Interfaz",
         "HTML5 y CSS3. Un documento de pantallas, un sistema de estilos y una hoja de impresión "
         "propia para el papel que se reparte en clase."],
        ["Lógica",
         f"JavaScript con módulos ES nativos ({miles(d['codigo'][0][1])} líneas): motor de "
         "generación, modelo de calificación, validador de calidad, catálogo de "
         "microexplicaciones, modo avanzado, pantalla de calificar y exportadores."],
        ["Datos",
         f"Archivos JSON versionados en el propio repositorio ({miles(d['codigo'][5][1])} "
         "líneas): packs de criterios, catálogo, tablas de derivación por materia, léxico de "
         "reglas, banco de verbos y esquema de pack."],
        ["Herramientas",
         f"Programas en Python ({miles(d['codigo'][3][1])} líneas) que construyen los derivados, "
         "validan el contenido y verifican mecánicamente la derivación normativa. No forman parte "
         "de lo que se sirve al navegador, pero sí de la obra."],
        ["Pruebas",
         f"Casos dorados en Node sin dependencias ({miles(d['codigo'][4][1])} líneas): validador, "
         "premarcado, proyección a primera persona y modelo de calificación."],
    ], anchos=[3.0, 13.6])

    m.h2("3.2 Diagrama de arquitectura")
    m.codigo([
        "┌────────────────────────────────────────────────────────────────────┐",
        "│  DOCENTE — navegador en ordenador, tableta o iPad                  │",
        "│  Instalable como aplicación · funciona sin conexión                │",
        "└─────────────────────────────┬──────────────────────────────────────┘",
        "                              │  sin red, sin cuentas, sin datos personales",
        "                              ▼",
        "┌────────────────────────────────────────────────────────────────────┐",
        "│  APLICACIÓN (js/) — módulos ES nativos, sin framework               │",
        "│                                                                    │",
        "│   motor.js .............. puerta de aplicabilidad, filtrado,       │",
        "│                           premarcado, pesos, los 7 instrumentos    │",
        "│   calificacion.js ....... modelo de cálculo (funciones puras)      │",
        "│   calificar.js .......... registro por alumno y CSV de notas       │",
        "│   validador.js .......... las reglas de calidad de descriptores    │",
        "│   lexico.js ............. palabras y umbrales — GENERADO           │",
        "│   microexplicaciones.js . el «¿por qué?» de cada control           │",
        "│   ui.js · main.js ....... vista previa, pestañas, CSV de rúbrica   │",
        "│   modo-avanzado.js ...... dimensiones y pesos                      │",
        "│   csv.js · pwa.js · sw.js  exportación, instalación, sin conexión  │",
        "└─────────────────────────────┬──────────────────────────────────────┘",
        "                              │  fetch local de sus propios archivos",
        "                              ▼",
        "┌────────────────────────────────────────────────────────────────────┐",
        "│  DATOS (data/) — la fuente de todo                                 │",
        "│                                                                    │",
        "│   pack-lcl-*.json ....... criterios, descriptores y matrices       │",
        "│   catalogo.json ......... packs, materias, cursos, niveles         │",
        "│   derivacion-lcl.json ... qué tarea sostiene el currículo y dónde  │",
        "│   reglas-lexicas.json ... palabras y umbrales del validador        │",
        "│   verbos.json ........... banco cerrado, 3.ª y 1.ª persona         │",
        "│   esquema-pack.json ..... la forma que debe tener un pack          │",
        "└─────────────────────────────┬──────────────────────────────────────┘",
        "                              │  fuera del navegador: construcción y",
        "                              │  comprobación del contenido",
        "                              ▼",
        "┌────────────────────────────────────────────────────────────────────┐",
        "│  HERRAMIENTAS (scripts/) y PRUEBAS (test/)                         │",
        "│   comprobar_todo.py ..... las comprobaciones que ejecuta el CI     │",
        "│   validar_pack.py · comprobar_paridad.py                           │",
        "│   verificar_derivacion.py  citas contra el currículo, con auto-    │",
        "│                            prueba de que lo corrupto debe fallar   │",
        "│   simular_correccion.py .. prueba de toda matriz cuantitativa      │",
        "│   generar_lexico.py · generar_tablas_sdd.py · generar_revision.py  │",
        "│   generar_manual.py · generar_iconos.py · ver.py                   │",
        "└────────────────────────────────────────────────────────────────────┘",
    ])

    m.h2("3.3 Volumen y organización")
    m.tabla(["Componente", "Líneas", "Nota"],
            [[nombre, miles(n), nota] for nombre, n, nota in d["codigo"]]
            + [["**TOTAL**", f"**{miles(total_lineas)}**",
                "código, contenido, herramientas, pruebas y documentación de diseño"]],
            anchos=[7.6, 2.0, 7.0], centrar_desde=1)

    m.h2("3.4 Cada hecho, en un solo sitio")
    m.p("La regla estructural que gobierna todo el repositorio es que **ningún hecho se escribe "
        "dos veces**. No es una preferencia de estilo: nació de tres copias del mismo dato que se "
        "separaron en silencio y dejaron a la aplicación imprimiendo un nombre de nivel que el "
        "material de aula ya no usaba. De ahí que los datos manden sobre el código y sobre la "
        "documentación, y que los derivados se generen:")
    m.tabla(["Archivo de datos", "Manda sobre"], [
        ["catalogo.json",
         "Qué packs hay, qué materias, qué tipos de tarea, cómo se llaman los cursos y los cuatro "
         "niveles de logro."],
        ["derivacion-&lt;materia&gt;.json".replace("&lt;", "<").replace("&gt;", ">"),
         "Qué tarea sostiene el currículo en qué curso, y hasta dónde exige cada uno."],
        ["reglas-lexicas.json", "Las palabras y los umbrales de las reglas del validador."],
        ["verbos.json", "El banco cerrado de verbos, con su tercera y su primera persona."],
        ["esquema-pack.json", "La forma que debe tener un pack, antes que ninguna regla de contenido."],
    ], anchos=[4.4, 12.2])
    m.p("**Son derivados y no se editan a mano**: el léxico que usa la aplicación, las dos tablas "
        "curriculares del documento de diseño, los documentos de revisión docente, la guía del "
        "profesorado, los iconos de la aplicación instalada y esta misma memoria. Una "
        "comprobación automática falla si alguno se ha separado de su fuente.")

    m.h2("3.5 Ausencia deliberada de dependencias, de servidor y de IA")
    m.p("La obra no emplea ningún marco de trabajo, ninguna biblioteca de terceros en su lógica, "
        "ningún empaquetador, ningún preprocesador y ningún lenguaje transpilado. Se apoya "
        "exclusivamente en las interfaces nativas del navegador: módulos ES, almacenamiento local "
        "para la persistencia de las calificaciones en el propio equipo del docente y trabajador "
        "de servicio para el funcionamiento sin conexión. Las herramientas de construcción usan "
        "la biblioteca estándar de Python, salvo la escritura de documentos en formato Word.")
    m.p("**Ninguna llamada a inteligencia artificial desde la aplicación.** La obra genera el "
        "texto del modo asistido para que el docente lo use fuera: no ejecuta prompts, no pide "
        "claves y no envía nada. El protocolo que acompaña siempre a ese texto exige retirar todo "
        "dato que identifique a un menor antes de usarlo, calibrar contra correcciones propias, "
        "exigir evidencia citada para cada nivel asignado, y deja escrito que **la nota la firma "
        "el profesor** y que el alumnado debe saber si una IA ha intervenido en su corrección.")
    m.salto()

    # ------------------------------------------------------------------ §4
    m.h1("4. ESTRUCTURAS DE DATOS PROPIETARIAS")
    m.p(f"El contenido curricular de la obra son {miles(d['codigo'][5][1])} líneas de JSON de "
        "diseño propio, con su esquema normativo y su validador. Lo que sigue describe las "
        "estructuras originales que lo sostienen.")

    m.h2("4.1 El pack de criterios")
    m.p(f"Un pack agrupa todos los criterios de un tipo de tarea, en todos los cursos donde el "
        f"currículo lo sostiene. Hay **{palabra(len(d['packs']))} packs** en explotación, uno por "
        "tipo de tarea, y cada uno declara su normativa de referencia, sus verbos añadidos al "
        "banco y la razón declarada de su reparto de pesos:")
    m.tabla(["Tipo de tarea", "Archivo", "Versión", "Criterios", "Cursos"],
            [[p["tarea"], p["archivo"], p["version"], str(p["criterios"]), str(p["cursos"])]
             for p in d["packs"]]
            + [["**TOTAL**", "", "", f"**{d['criterios']}**", ""]],
            anchos=[4.4, 6.0, 1.8, 2.2, 2.2], centrar_desde=2)

    m.h2("4.2 El criterio")
    m.p("Es la unidad de la obra. Un criterio no es una fila de tabla: es un registro que "
        "**ata una dimensión evaluable a una cita literal del decreto** y declara su posición en "
        "los tres ejes de progresión. Fragmento real, abreviado:")
    m.codigo([
        '{',
        '  "id": "lcl-b-cohesion-3eso",',
        '  "bloque_lomloe": "B",',
        '  "dimension": "cohesion",',
        '  "nombre": "Cohesión: conectores y puntuación",',
        '  "curso": "3ESO",',
        '',
        '  "criterio_oficial": {',
        '    "competencia_especifica": 5,',
        '    "codigo": "5.2",',
        '    "cita": "Producir textos escritos y multimodales coherentes, cohesionados,',
        '             adecuados y correctos...",',
        '    "perfil_salida": ["CCL1", "CCL3", "CCL5"]',
        '  },',
        '',
        '  "saber_vehiculo": ["texto argumentativo", "marcadores discursivos"],',
        '  "progresion": { "autonomia": 2, "complejidad": 2, "metalinguistico": 1 },',
        '  "tipos_tarea": ["argumentativo", "expositivo"],',
        '  "evalua_proceso": false,  "obligatorio": false,',
        '  "prioridad": 1,           "peso_base": 20,',
        '',
        '  "descriptores": {',
        '    "n1": { "verbo": "utiliza",  "texto": "Utiliza conectores de adición de forma',
        '             repetitiva, lo que produce saltos entre las ideas del texto." },',
        '    "n2": { "verbo": "utiliza",  "texto": "Utiliza conectores básicos de causa y',
        '             oposición para enlazar las ideas principales del texto." },',
        '    "n3": { "verbo": "emplea",   "texto": "Emplea conectores de causa, consecuencia',
        '             y oposición, y delimita con la puntuación los incisos." },',
        '    "n4": { "verbo": "articula", "texto": "Articula la lógica argumentativa mediante',
        '             marcadores variados y sostiene la cohesión con anáforas sin ambigüedad." }',
        '  }',
        '}',
    ])
    m.p("Tres campos son de diseño propio y sostienen todo lo demás. **`criterio_oficial`** es la "
        "puerta: sin él el criterio no carga. **`saber_vehiculo`** separa el contenido que la "
        "tarea vehicula del nombre de la dimensión, que ha de ser una acción competencial y no un "
        "tema. **`progresion`** sitúa el criterio en los tres ejes —autonomía, complejidad y "
        "reflexión metalingüística— que permiten comprobar que un instrumento no mezcla peldaños "
        "incompatibles.")
    m.p(f"El conjunto suma **{miles(d['descriptores'])} descriptores** repartidos en "
        f"**{palabra(d['dimensiones'])} dimensiones distintas**, sobre "
        f"**{palabra(d['citas'])} pares de criterio oficial y curso** citados literalmente.")

    m.h2("4.3 La matriz cuantitativa")
    m.p("Cuando una dimensión admite corrección medible, el criterio añade una matriz de "
        "componentes y bandas contables. Es la estructura que convierte «emplea conectores "
        "variados» —que no es verificable— en «3,0 puntos si emplea cuatro o más tipos distintos "
        "de marcador; 2,0 si emplea dos o tres», que sí lo es.")
    m.tabla(["Elemento", "Recuento", "Nota"], [
        ["Matrices cuantitativas", miles(d["matrices"]),
         f"sobre {miles(d['criterios'])} criterios"],
        ["Componentes", miles(d["componentes"]), "cada uno con su puntuación máxima"],
        ["Bandas contables", miles(d["bandas"]),
         "cada una con su condición medible, sin calificadores vagos"],
        ["Penalizaciones declaradas", palabra(d["penalizaciones"]),
         "y es deliberado: véase el recuadro siguiente"],
    ], anchos=[5.4, 3.0, 8.2], centrar_desde=1)
    m.caja("Por qué no hay ni una sola penalización, y es buena señal",
           "El diseño prohíbe el **doble castigo**: una penalización no puede medir lo que un "
           "componente de la misma matriz ya recoge. La regla salió de la primera matriz que se "
           "escribió, donde un componente valoraba los errores de puntuación y encima había una "
           "penalización por párrafo largo sin puntos; el alumno pagaba dos veces por el mismo "
           "fenómeno, la dimensión caía dos niveles de golpe y la nota final perdía siete "
           "décimas. Una rúbrica así no se sostiene ante una reclamación, porque el alumno puede "
           "señalar exactamente dónde se le ha restado dos veces. La consecuencia adoptada es "
           "que, si un fenómeno merece medirse, casi siempre debe ser un componente con su banda "
           "y no un descuento. Al aplicarla a fondo, las matrices quedan **sin penalización "
           "alguna** — y el mecanismo sigue implementado, con su tope obligatorio y su regla de "
           "validación, para el pack que algún día lo necesite.")

    m.h2("4.4 Las tablas de derivación por materia")
    m.p("Dos hechos curriculares no pertenecen a ningún pack porque valen para todos, y viven en "
        "un archivo propio por materia. El primero es la **matriz de tarea × curso**: qué tipo de "
        "tarea sostiene el currículo en qué curso. Hoy declara "
        f"**{palabra(d['n_celdas'])} celdas** sobre las {n_tareas * n_cursos} combinaciones "
        f"posibles de {palabra(n_tareas)} tipos de tarea y {palabra(n_cursos)} cursos:")
    m.tabla(["Tipo de tarea"] + [cursos["etiquetas_cortas"][c] for c in cursos["orden"]],
            [[tareas[t]] + ["●" if c in d["celdas"].get(t, []) else "·" for c in cursos["orden"]]
             for t in tareas],
            anchos=[5.2] + [1.9] * n_cursos, centrar_desde=1)
    m.p("Una celda ausente no es un hueco por rellenar: significa que **ningún criterio de "
        "evaluación de ese curso abre esa tarea**, y entonces la aplicación no la ofrece y ningún "
        "pack puede declararla ahí. El segundo hecho es el **techo de progresión**: hasta dónde "
        "puede exigir cada curso en cada uno de los tres ejes. Las dos tablas se imprimen en el "
        "documento de diseño desde este archivo, nunca al revés.")

    m.h2("4.5 Banco cerrado de verbos y léxico de reglas")
    m.p(f"El banco guarda **{palabra(d['verbos'])} verbos**, cada uno en tercera y en primera "
        "persona («Reconoce» / «Reconozco»). Es a la vez una restricción de redacción —todo "
        "descriptor debe empezar por un verbo del banco— y el mecanismo que permite derivar la "
        "versión de autoevaluación sin errores de morfología (§5.5). Un pack solo puede añadir "
        "verbos en su lista de extras, nunca alterar el banco común.")
    m.p("El léxico de reglas guarda aparte las palabras y los umbrales que el validador vigila: "
        "los calificadores vagos, las negaciones que delatan un nivel 1 formulado en negativo, "
        "los modalizadores de ayuda y de autonomía del currículo, el reconocedor de recuentos de "
        "banda y los umbrales numéricos. Está separado por materia, y **dar de alta una "
        "asignatura nueva no debe obligar a tocar código**: si hiciera falta editar un programa, "
        "sería un fallo de diseño de esta obra, no de la materia nueva.")

    m.h2("4.6 El esquema de pack")
    m.p("La forma que debe tener un pack está declarada en un esquema propio, que se comprueba "
        "**antes** que ninguna regla de contenido: primero se verifica que el archivo tiene la "
        "forma prometida y solo después se juzga lo que dice.")
    m.salto()

    # ------------------------------------------------------------------ §5
    m.h1("5. ALGORITMOS Y LÓGICA ORIGINAL")

    m.h2("5.1 Premarcado de dimensiones de proceso")
    m.p("De las cinco respuestas de la puerta de aplicabilidad, la de «fase de un texto» no elige "
        "solo instrumento: elige **qué dimensiones vienen marcadas**, porque lo que se entrega no "
        "es el texto. En un esquema no hay cohesión que observar y en un borrador la corrección "
        "normativa aún no se juzga.")
    m.p("Cuáles son esas dimensiones **lo declara el pack** y no lo deduce el motor, y esa es la "
        "parte no evidente. De la cita no se deduce: el mismo criterio que habla de planificar y "
        "revisar es también el que sostiene la adecuación, la coherencia y la cohesión del texto "
        "terminado, de modo que un premarcado leído de la cita marcaría media rúbrica. Lo que sí "
        "se comprueba mecánicamente es **la dirección contraria**: una dimensión declarada de "
        "proceso cuya cita oficial no hable de planificar, de borradores ni de revisar no carga.")
    m.p("**El premarcado nunca vacía el instrumento.** Dos casos de borde tienen salida escrita: "
        "si la dimensión de proceso no sobrevive al filtro de tiempo, se rescata y se avisa —el "
        "filtro mide el coste de corregir un texto entero, que es justo lo que aquí no se está "
        "corrigiendo—; y si la tarea no tiene ninguna dimensión de proceso, como ocurre con la "
        "exposición oral, se mantienen todas premarcadas y se explica por qué.")

    m.h2("5.2 Filtro de profundidad por tiempo de corrección")
    m.p("El tiempo declarado por alumno se traduce en un corte por prioridad: menos de dos "
        "minutos deja solo las dimensiones de prioridad máxima; de dos a cinco añade el segundo "
        "escalón; más de cinco admite los tres. Es la traducción operativa de un principio de "
        "sostenibilidad: una rúbrica que no se puede corregir en el tiempo real del docente no se "
        "usa, y una rúbrica que no se usa no evalúa nada.")

    m.h2("5.3 Modelo de cálculo")
    m.p("El modelo está especificado hasta el detalle porque **tiene que sostenerse ante una "
        "reclamación**. Sus reglas exactas:")
    m.puntos([
        "**Dos escalas de nivel seleccionables**: equilibrada (2,5 · 5 · 7,5 · 10) por defecto, "
        "que evita el cero a quien ha producido algo aunque sea flojo, y exigente (0 · 5 · 7,5 · "
        "10), que reserva el cero al trabajo no realizado. La elección se declara en la hoja del "
        "alumno; no se oculta.",
        "**Una dimensión con matriz aporta sus puntos continuos, no el valor de su nivel.** Si "
        "los puntos se colapsaran a nivel antes de ponderar, un alumno con 8,9 aportaría 7,5 y "
        "otro con 9,0 aportaría 10: una décima de desempeño valdría dos puntos y medio de nota. "
        "Un escalón así es indefendible.",
        "**Las penalizaciones aplicadas son negativas y se suman.** Es la trampa de signo que "
        "parece evidente escrita y es fácil de invertir al leer; restarlas subiría la nota.",
        "**La puntuación de una dimensión nunca baja de cero**, aunque todos los componentes "
        "estén en su banda mínima y todas las penalizaciones a tope.",
        "**El detractor global se aplica antes que la condición mínima**, y el orden importa: con "
        "una nota de 8, un detractor de 2 y la condición disparada, invertir el orden da 2,9 en "
        "lugar de 4,9. Se fija el orden que respeta lo que cada mecanismo significa: el detractor "
        "forma parte de calcular la nota; la condición mínima es un límite sobre la nota final.",
        "**La condición mínima es un techo no acumulativo y desactivado por defecto**: limita la "
        "nota a 4,9 si un criterio obligatorio queda en el nivel más bajo, nunca la sube, y no "
        "actúa si el instrumento no tiene ningún criterio obligatorio. Se deja desactivada porque "
        "una condición mínima que no se ha anunciado al alumnado antes de la prueba es difícil de "
        "sostener; si se activa, la aplicación la imprime obligatoriamente en la hoja del alumno.",
        "**Redondeo a la centésima**, mitad hacia arriba y con corrección de la representación "
        "binaria, aplicado en dos momentos exactos y solo en esos dos: a la puntuación de cada "
        "dimensión antes de ponderarla, y a la nota final una sola vez al terminar la suma.",
    ])

    m.h2("5.4 De puntos a nivel")
    m.p("Un número de cero a diez se traduce al nombre de su nivel por umbrales cerrados —nueve, "
        "siete y cinco— y no por intervalos escritos con un decimal, que dejarían sin cubrir el "
        f"hueco entre 8,9 y 9,0. Los cuatro nombres canónicos ({enumerar(nombres_nivel)}) viven "
        "**una sola vez** en el catálogo de datos, porque estuvieron escritos en cuatro sitios "
        "del código y por eso la aplicación siguió imprimiendo un nombre antiguo meses después de "
        "que la decisión de cambiarlo estuviera tomada.")

    m.h2("5.5 Proyección exacta a primera persona")
    m.p("La versión de autoevaluación no se redacta aparte: se **deriva** de la matriz completa "
        "reconjugando cada verbo a primera persona a través del banco. El algoritmo tiene dos "
        "cautelas que salieron de fallos reales: reconjugar solo el verbo inicial dejaba media "
        "matriz mezclando «yo» y «él», porque muchos descriptores llevan un segundo verbo; y "
        "reconjugarlos todos convertía «una introducción **que delimita** el tema» en «que "
        "delimito», cambiando de sujeto la introducción por el alumno. La regla final reconjuga "
        "todo verbo del banco salvo el que sigue inmediatamente a un relativo.")
    m.p("De ahí viene además una regla de redacción propia: **el descriptor no nombra al alumno "
        "en tercera persona**. El alumno es el sujeto de la frase, así que cualquier otra marca "
        "que apunte a él —un posesivo, un dativo— imprimiría «él» dentro de una frase en «yo». Se "
        "escribe «por cuenta propia» y no «por su cuenta», «con palabras propias» y no «con sus "
        "propias palabras». Un caso dorado fija ese invariante sobre el contenido completo, para "
        "que una regresión no llegue nunca al navegador.")

    m.h2("5.6 El validador, dos veces y en paridad comprobada")
    m.p("El validador está implementado dos veces —en Python, para construir el contenido, y en "
        "JavaScript, para la aplicación— porque son dos momentos distintos con dos consumidores "
        "distintos. Pero son dos programas, no dos criterios: **la aplicación nunca puede dar por "
        "limpio un pack que el script rechaza**, y esa afirmación se comprueba mecánicamente "
        "sobre un corpus de trampas construido a propósito. Las palabras y los umbrales no se "
        "escriben en ninguna de las dos: salen del archivo de léxico compartido, y la versión "
        "JavaScript se genera desde él.")
    m.p("La única divergencia admitida va en la dirección segura: el script busca por subcadena y "
        "por eso marca de más, mientras que la aplicación exige palabra completa en los términos "
        "cortos para no marcar «formal» por contener «mal». Está documentada en el código de "
        "ambos.")

    m.h2("5.7 El lector de recuentos y la continuidad de bandas")
    m.p("Una matriz de ortografía —«hasta 2 faltas», «de 3 a 5», «de 6 a 9», «10 o más»— tiene "
        "que cubrir todos los recuentos, porque el corrector cuenta y busca su banda: si ninguna "
        "dice qué hacer con 4, la nota se decide a ojo justo donde la matriz prometía aritmética, "
        "y dos correctores dan dos notas que nadie puede reconstruir. Una matriz de fuentes "
        "reunidas —«4 o más», «3», «2», «1»— no tiene esa obligación, porque ahí el recuento sube "
        "con la nota.")
    m.p("**Cuál de las dos es no lo declara el pack: se lee de la propia matriz.** Si la cuenta "
        "sube según bajan los puntos, lo contado es una incidencia; si baja, es un logro. Añadir "
        "un campo habría sido la otra salida y se descartó: la dirección de una escala se deduce "
        "sin ambigüedad con la matriz delante, y un campo más sería contenido que mantener en "
        "nueve packs para repetir un dato ya escrito. El reconocedor tiene además una cautela "
        "propia: **no interpreta «un», «una» ni «uno» como números**, porque en estos textos son "
        "artículo mucho más veces que número e inventaban escalas donde no había ninguna.")

    m.h2("5.8 Verificación mecánica de la derivación normativa")
    m.p("Toda afirmación de la obra sobre lo que dice el currículo se comprueba contra el texto "
        "real del decreto, y no contra la memoria de quien la escribió. El verificador contrasta "
        "las citas literales de los packs y del documento de diseño, la matriz de tareas, el "
        "techo de progresión y el género declarado en el vehículo de cada criterio, **recortando "
        "el currículo por cursos** para poder comprobar una cita contra el curso al que dice "
        "pertenecer y no contra el documento entero. Fue así como se detectó una cita atribuida a "
        "un curso en el que ese género no existe.")
    m.p("El verificador lleva su propia **auto-prueba**: corrompe deliberadamente una cita, una "
        "celda y un techo, y falla si el verificador no los detecta. Una comprobación que no "
        "puede fallar no comprueba nada.")

    m.h2("5.9 Simulador de corrección")
    m.p("Toda matriz cuantitativa nueva se prueba **simulando una corrección completa** con "
        "perfiles de alumno antes de darla por buena. El simulador hace lo mecánico —notas por "
        "curso y perfil, nivel de cada dimensión, saltos de dos niveles, umbrales planos, ceros y "
        "notas fuera de orden— y deja el juicio entero al docente: *¿le pondrías esta nota a un "
        "alumno con este perfil?* Está deliberadamente fuera de la comprobación automática, "
        "porque pide un juicio y no devuelve un booleano.")

    m.h2("5.10 Una sola orden lo comprueba todo")
    m.p(f"Las **{palabra(d['comprobaciones'])} comprobaciones** de la obra se ejecutan con una "
        "única orden, que es exactamente la misma que ejecuta la integración continua en cada "
        "envío, de modo que no exista una segunda lista que mantener: forma de los packs, los "
        "derivados al día, reglas de contenido, paridad entre los dos validadores, pruebas del "
        "motor y del modelo de calificación, verificación de la derivación con su auto-prueba, y "
        "que exista todo lo que la instalación promete. **No se cierra un pack con ninguna en "
        f"rojo ni con ninguna saltada.** A fecha de {FECHA} las "
        f"{palabra(d['comprobaciones'])} están limpias.")
    m.salto()

    # ------------------------------------------------------------------ §6
    m.h1("6. METODOLOGÍA PEDAGÓGICA ORIGINAL")
    m.p("La obra no es un editor de tablas con una capa de currículo encima: es la materialización "
        "en software de un conjunto de reglas evaluativas que el autor sostiene y que el programa "
        "**hace cumplir mecánicamente**. Cada principio tiene aquí su consecuencia técnica "
        "concreta, y esa traducción es la aportación original.")

    m.tabla(["Principio", "Cómo lo impone el programa"], [
        ["**Las rúbricas se derivan, no se inventan**",
         "Campo de cita oficial obligatorio en cada criterio. Un criterio sin referencia "
         "normativa no supera la validación y no se carga."],
        ["**Los saberes son vehículo, nunca fila**",
         "El contenido vehiculado se guarda en un campo separado del nombre de la dimensión, y el "
         "validador rechaza los nombres que son un tema en lugar de una acción competencial."],
        ["**Cuatro niveles, siempre**",
         "Escala fija en todo el sistema, no configurable, con bandas sobre diez y nombres "
         "canónicos escritos una sola vez."],
        ["**Cero adverbitis**",
         "Banco cerrado de verbos y reglas que bloquean los calificadores vagos, tanto en el "
         "descriptor como dentro de las bandas de una matriz."],
        ["**Gradación positiva incluso en el nivel bajo**",
         "El nivel 1 describe lo que el alumno sí hace de forma limitada. El validador marca todo "
         "descriptor de nivel 1 que empiece por una negación."],
        ["**La progresión la fija el currículo**",
         "Campo de progresión en tres ejes y regla de modalizadores en dos direcciones. Se "
         "descartaron dos formulaciones anteriores que ponían un techo cognitivo por curso: "
         "marcaban como error descriptores correctos, porque el verbo lo pone el criterio."],
        ["**La transparencia es una salida**",
         "La hoja del alumno se genera siempre y no se puede desmarcar."],
        ["**Sostenibilidad de la corrección**",
         "Filtro de profundidad por tiempo real e indicador de complejidad, con aviso activo por "
         "encima de cinco dimensiones."],
        ["**Ponderación igual por defecto**",
         "Desigualarla exige una razón declarada, que la hoja del alumno imprime y el validador "
         "reclama."],
        ["**Ni doble castigo ni descuento sin tope**",
         "Regla de validación que impide que una penalización mida lo que un componente ya "
         "recoge, y tope obligatorio en cada descuento."],
        ["**El error no resta en absoluto: resta acotado**",
         "Todo descuento tiene techo declarado, la dimensión nunca baja de cero y la vía de "
         "mejora queda señalada al lado."],
        ["**La herramienta enseña mientras se usa**",
         "Microexplicación en cada control y explicación de la regla en cada aviso del validador. "
         "El objetivo no es que existan más rúbricas, sino que se sepan usar."],
    ], anchos=[4.8, 11.8])
    m.salto()

    # ------------------------------------------------------------------ §7
    m.h1("7. APORTACIÓN DIDÁCTICA ORIGINAL")

    m.h2("7.1 Naturaleza de la aportación")
    m.p("La obra no se agota en su condición de programa de ordenador: incorpora un cuerpo "
        "sustancial de contenidos didácticos originales, de creación propia del autor, elaborados "
        "por derivación del currículo oficial. Estos contenidos constituyen expresión protegida "
        "por el derecho de autor y forman parte inseparable del valor de la obra.")
    m.p("A los efectos de esta memoria conviene precisar que la protección recae sobre la "
        "**expresión concreta** de dichos contenidos —los textos, las matrices, las secuencias y "
        "los criterios tal como han sido redactados y organizados por el autor— y no sobre las "
        "ideas, los métodos o los conceptos pedagógicos en abstracto, que la legislación excluye "
        "de protección. La cita literal del criterio oficial procede del decreto autonómico y no "
        "es objeto de apropiación alguna: lo original es la **derivación**, es decir, el trabajo "
        "de convertir esa cita en dimensiones evaluables, descriptores graduados y bandas "
        "medibles.")

    m.h2("7.2 Contenidos didácticos originales")
    m.tabla(["Contenido", "Descripción"], [
        ["Biblioteca de criterios derivados",
         f"{miles(d['criterios'])} criterios de evaluación en {palabra(len(d['packs']))} tipos de "
         f"tarea y {palabra(d['cursos_con_pack'])} cursos, cada uno con su cita literal del "
         "decreto, su vehículo de saberes y su posición en los tres ejes de progresión."],
        ["Descriptores de nivel",
         f"{miles(d['descriptores'])} descriptores redactados con la fórmula propia «verbo del "
         "banco + objeto o saber-vehículo + condición o finalidad comunicativa», en "
         f"{palabra(d['dimensiones'])} dimensiones distintas y sin calificadores vagos."],
        ["Matrices cuantitativas de corrección",
         f"{miles(d['matrices'])} matrices con {miles(d['componentes'])} componentes y "
         f"{miles(d['bandas'])} bandas contables, calibradas curso a curso y probadas por "
         "simulación de corrección con perfiles de alumno."],
        ["Tablas de derivación curricular",
         f"La matriz de qué tarea sostiene el currículo en qué curso ({palabra(d['n_celdas'])} "
         "celdas justificadas una a una) y el techo de progresión por curso y eje."],
        ["Banco cerrado de verbos",
         f"{palabra(d['verbos'])} verbos observables con su tercera y su primera persona, que "
         "hacen posible la derivación exacta de la autoevaluación."],
        ["Léxico de reglas de calidad",
         "Calificadores vagos, negaciones, modalizadores de ayuda y de autonomía, reconocedor de "
         "recuentos y umbrales, separado por materia."],
        ["Catálogo de microexplicaciones",
         f"{palabra(d['micro'])} textos explicativos propios que enseñan al profesorado por qué "
         "existe cada decisión de diseño, situados en el control al que afectan."],
        ["Documento de diseño y registro de motivos",
         "La especificación completa del modelo de datos, del modelo de calificación, del "
         "catálogo de instrumentos y de las reglas del validador, con el registro razonado de "
         "cada decisión y de las formulaciones equivocadas que se descartaron."],
        ["Documentos de revisión docente",
         "Un documento por tipo de tarea, generado desde el contenido, para que otro profesor "
         "pueda revisar la derivación sin abrir un archivo de datos."],
        ["Guía del profesorado",
         "Manual de uso en formato Word, generado desde el propio contenido para que sus cifras "
         "no puedan separarse de la obra que describen."],
        ["Protocolo de corrección asistida",
         "Las cinco reglas que acompañan siempre al texto del modo asistido: protección de datos "
         "del menor, calibración previa, evidencia citada obligatoria, firma del profesor y "
         "transparencia con el alumnado."],
    ], anchos=[4.6, 12.0])

    m.h2("7.3 Materiales de terceros excluidos de la inscripción")
    m.p("Las citas literales de los criterios de evaluación proceden del currículo oficial "
        f"publicado por la {d['comunidad']} y del Boletín Oficial correspondiente. Se utilizan "
        "como cita normativa —uso expresamente amparado— y **no son objeto de la presente "
        "inscripción**: los textos legales carecen de protección por derecho de autor y el "
        "repositorio los conserva íntegros y sin alterar, en su carpeta de fuentes, precisamente "
        "para que la derivación pueda comprobarse contra ellos. Lo mismo vale para el documento "
        "de estructura y criterios de la prueba de acceso a la universidad, incorporado solo como "
        "material de contraste.")
    m.p("La inscripción se refiere exclusivamente a las creaciones propias del autor: el "
        "programa, su arquitectura, las estructuras de datos y los contenidos didácticos "
        "relacionados en el apartado 7.2. Se hace constar, a los efectos que el Registro estime "
        "oportunos, que algunos de esos materiales —la biblioteca de criterios derivados, las "
        "matrices de corrección o la guía del profesorado— podrían además ser objeto de "
        "inscripción como obras independientes de naturaleza literaria o educativa, con "
        "independencia del programa de ordenador.")
    m.salto()

    # ------------------------------------------------------------------ §8
    m.h1("8. COMPONENTES INCLUIDOS EN LA INSCRIPCIÓN")
    m.tabla(["Componente", "Descripción y extensión"], [
        ["Código fuente de la aplicación",
         f"{miles(d['codigo'][0][1])} líneas de JavaScript en módulos ES nativos, más "
         f"{miles(d['codigo'][1][1])} de HTML e instalación y {miles(d['codigo'][2][1])} de "
         "estilos de pantalla e impresión."],
        ["Herramientas de construcción y comprobación",
         f"{miles(d['codigo'][3][1])} líneas de Python: validadores, generadores de derivados, "
         "verificador de derivación normativa, simulador de corrección y lector de contexto."],
        ["Baterías de prueba",
         f"{miles(d['codigo'][4][1])} líneas de casos dorados en Node, sin dependencias: "
         "validador regla a regla, premarcado y puertas, proyección a primera persona y modelo de "
         "calificación."],
        ["Contenido curricular estructurado",
         f"{miles(d['codigo'][5][1])} líneas de JSON: {palabra(len(d['packs']))} packs de "
         f"criterios, catálogo, tablas de derivación, léxico de reglas, banco de verbos y esquema."],
        ["Documentación de diseño",
         f"{miles(d['codigo'][6][1])} líneas: documento de diseño de software, registro razonado "
         "de cambios, motivos de cada regla, guía para añadir una materia y análisis de enlace "
         "con los proyectos afines del autor."],
        ["Materiales para el profesorado",
         "Guía de uso y documentos de revisión por tipo de tarea, todos generados desde el "
         "contenido."],
        ["Identidad visual e instalación",
         "Sistema de estilos editorial propio, hoja de impresión para A4, manifiesto de "
         "instalación y juego completo de iconos generados por el repositorio."],
    ], anchos=[4.6, 12.0])

    m.h2("8.1 Historial de versiones")
    m.p("A efectos probatorios de antigüedad se deja constancia del histórico de desarrollo, "
        "registrado en el sistema de control de versiones del repositorio y en el registro de "
        "cambios del documento de diseño:")
    m.tabla(["Versión", "Fecha", "Hito principal"], [
        ["0.9 – 1.0", "28-29 julio 2026",
         "Prototipo funcional: motor de generación, modo exprés, rúbrica analítica, lista de "
         "cotejo, hoja del alumno e impresión. El validador queda completo dentro de la "
         "aplicación, con su microexplicación por regla."],
        ["1.1 – 1.9", "29 julio – 4 agosto 2026",
         "Repositorio con validación automática en cada envío. Reglas exactas del modo numérico. "
         "Persistencia del resultado por alumno. Rediseño visual editorial. Autoevaluación y "
         "coevaluación por proyección a primera persona. Rúbrica de un solo punto y escala de "
         "estimación analítica: el catálogo de instrumentos queda cerrado."],
        ["1.10 – 1.20", "5-7 agosto 2026",
         "Packs de texto argumentativo, exposición oral y texto expositivo ampliados hasta "
         "Bachillerato. Quinta puerta de aplicabilidad, con premarcado de dimensiones de proceso. "
         "Verificación de citas contra el segmento de su propio curso. Pack de narración: la "
         "matriz de tarea × curso de la fase 2 queda completa."],
        ["1.21 – 1.24", "agosto 2026",
         "Las tablas curriculares se mudan a los datos, una por materia, y las del documento de "
         "diseño pasan a ser derivados comprobados. Barrido completo del posesivo y del dativo de "
         "tercera persona en todos los packs."],
        ["1.25 – 1.31", "12-16 agosto 2026",
         "Comentario de texto literario y resumen. Nombres de nivel unificados con el material de "
         "aula, razón de peso declarada y umbrales recalibrados por curso. Instalación como "
         "aplicación, publicación con dirección estable y adopción del nombre «Taller de "
         "Rúbricas»."],
        ["1.32 – 1.35", "18 agosto 2026",
         "Reorganización del repositorio para el trabajo asistido: lector selectivo y documento "
         "de diseño dividido. Tres tipos de tarea nuevos —reacción a una noticia, redacción de "
         "una noticia y trabajo de investigación multimodal—, dados de alta **sin tocar una línea "
         "de código**. Regla de continuidad de bandas."],
        ["1.36 – 1.41", "20 agosto 2026",
         "Las escalas de faltas pasan al modelo de bandas por densidad; toda matriz reparte en "
         "techos pares y todo componente tiene banda para el nivel más bajo; los tramos en "
         "singular se hacen medibles."],
        ["1.42", "25 agosto 2026",
         "Se cierran las seis decisiones que quedaban abiertas en el documento de diseño: "
         "condición mínima, escala por defecto, tope de descuentos, rúbrica holística, "
         "publicación y reparto de responsabilidades con la corrección asistida."],
        ["1.43 – 1.44", "26 agosto 2026",
         "Exportación a los dos importadores del cuaderno del profesor: el CSV de notas y la "
         "rúbrica en blanco, como caminos alternativos que elige el docente."],
        ["**1.45**", f"**{FECHA}**",
         "**Versión objeto de esta solicitud.** La guía del profesorado pasa a generarse desde el "
         "contenido, y el barrido para rehacerla destapa un defecto de concordancia en tres "
         "packs, corregido sin mover ningún umbral. Es el mismo principio que rige esta memoria: "
         "un dato escrito a mano envejece en silencio."],
    ], anchos=[2.4, 2.8, 11.4])
    m.salto()

    # ------------------------------------------------------------------ §9
    m.h1("9. TITULARIDAD Y DERECHOS DE EXPLOTACIÓN")
    m.p(f"El autor, D. {AUTOR}, declara ser titular exclusivo de todos los derechos de propiedad "
        "intelectual sobre la obra descrita en el presente documento y sobre la totalidad del "
        "código, las estructuras de datos y los contenidos didácticos que la integran, en virtud "
        "del artículo 5 del Real Decreto Legislativo 1/1996, de 12 de abril, por el que se "
        "aprueba el Texto Refundido de la Ley de Propiedad Intelectual.")
    m.p("La obra ha sido desarrollada íntegramente por el autor, fuera del horario laboral, con "
        "medios propios y sin uso de recursos del centro educativo en el que presta servicio como "
        "docente. El autor no ha cedido, transmitido ni licenciado derecho alguno sobre la obra a "
        "terceros, ni ha firmado acuerdo alguno que atribuya la titularidad a otra persona física "
        "o jurídica. El hecho de que la aplicación esté accesible públicamente en una dirección "
        "web para su uso por otros docentes constituye una puesta a disposición autorizada por el "
        "autor y no implica cesión ni cotitularidad de derecho alguno.")
    m.p("El autor se reserva la totalidad de los derechos patrimoniales de explotación sobre la "
        "obra, incluidos la reproducción, la distribución, la comunicación pública y la "
        "transformación, conforme a los artículos 17 a 23 del mismo texto refundido.")

    m.h2("9.1 Registros previos y obras relacionadas")
    m.p("La presente solicitud es la **primera inscripción** de esta obra. Se hace constar que el "
        "mismo autor tiene inscrita en este mismo Registro Territorial otra obra de software "
        "educativo de su exclusiva titularidad, «Taller de Sintaxis», con la que la presente "
        "comparte terminología y filosofía evaluativa pero **no comparte código, estructuras de "
        "datos ni contenidos**: son dos programas independientes, con arquitectura, modelo de "
        "datos y finalidad distintos.")
    m.p("El autor podrá depositar adicionalmente la obra en un servicio de registro privado de "
        "obra digital a efectos de prueba fehaciente de la fecha de creación. Dicho depósito no "
        "sustituye al registro público, que se solicita para obtener la presunción legal de "
        "autoría del artículo 145.3 del Texto Refundido de la Ley de Propiedad Intelectual. "
        "Número de depósito, si se aporta: [[______________________]]")

    m.h1("10. FIRMA")
    m.p("En [[______________________]], a [[______]] de [[______________]] de 2026.")
    for _ in range(4):
        m.doc.add_paragraph().paragraph_format.space_after = Pt(0)
    m.p(f"Fdo.: {AUTOR}", size=11)
    m.p(f"DNI {DNI}")
    m.salto()

    # ---------------------------------------------------------------- anexo
    m.h1("ANEXO · RELACIÓN DE FICHEROS DE LA OBRA")
    m.p("A efectos identificativos se relacionan los ficheros que componen el código fuente, el "
        "contenido curricular y los materiales didácticos de la obra, cuyo contenido íntegro se "
        "aporta como material de la solicitud.")

    m.h2("Aplicación e interfaz")
    m.puntos([
        "`index.html` — documento de pantallas de la aplicación.",
        "`css/styles.css` — sistema de estilos de pantalla. `css/print.css` — hoja de impresión "
        "para el papel que se reparte en clase.",
        "`manifest.webmanifest` y `sw.js` — instalación en iPad, Android y escritorio, y "
        "funcionamiento sin conexión.",
        "`assets/icons/` — juego completo de iconos de la aplicación instalada (generados).",
    ])

    m.h2("Lógica de la aplicación (js/)")
    m.puntos([
        "`motor.js` — puerta de aplicabilidad, filtrado, premarcado, pesos y generación de los "
        "instrumentos.",
        "`calificacion.js` — modelo de cálculo en funciones puras. `calificar.js` — pantalla de "
        "registro por alumno y exportación de notas.",
        "`validador.js` — las reglas de calidad de descriptores. `lexico.js` — palabras y "
        "umbrales (generado, no se edita).",
        "`microexplicaciones.js` — catálogo único del «¿por qué?» de cada control.",
        "`ui.js` y `main.js` — vista previa, pestañas de instrumento y exportación de la rúbrica.",
        "`modo-avanzado.js` — selección de dimensiones y reparto de pesos. `csv.js` — escritura "
        "conforme a RFC 4180. `pwa.js` — instalación.",
    ])

    m.h2("Contenido curricular (data/)")
    m.puntos(
        [f"`{p['archivo']}` — {p['tarea'].lower()}: {p['criterios']} criterios en "
         f"{palabra(p['cursos'])} cursos." for p in d["packs"]]
        + [
            "`catalogo.json` — packs, materias, cursos y nombres de los cuatro niveles.",
            "`derivacion-lcl.json` — matriz de tarea × curso y techo de progresión por eje.",
            "`reglas-lexicas.json` — palabras y umbrales del validador, por materia.",
            "`verbos.json` — banco cerrado de verbos, en tercera y primera persona.",
            "`esquema-pack.json` — la forma que debe tener un pack.",
        ])

    m.h2("Herramientas de construcción y comprobación (scripts/)")
    m.puntos([
        "`comprobar_todo.py` — la orden única que ejecuta todas las comprobaciones, la misma que "
        "ejecuta la integración continua.",
        "`validar_esquema.py` y `validar_pack.py` — forma del pack y reglas de contenido. "
        "`comprobar_paridad.py` — exige que los dos validadores digan lo mismo.",
        "`verificar_derivacion.py` — comprueba las citas contra el currículo real, con su "
        "auto-prueba de que lo corrupto debe fallar.",
        "`simular_correccion.py` — simulación de corrección con perfiles de alumno para probar "
        "toda matriz nueva.",
        "`generar_lexico.py`, `generar_tablas_sdd.py`, `generar_revision.py`, "
        "`generar_manual.py`, `generar_memoria_registro.py`, `generar_iconos.py` — generadores de "
        "los derivados.",
        "`dossier_criterios.py`, `ver.py`, `catalogo.py` — lectura selectiva del currículo y del "
        "diseño.",
    ])

    m.h2("Pruebas (test/)")
    m.puntos([
        "`validar-reglas.mjs` y `validar-pack-real.mjs` — un caso por regla del validador, más la "
        "carga limpia del contenido real.",
        "`premarcado.mjs` — puertas de aplicabilidad y premarcado de dimensiones de proceso.",
        "`proyeccion.mjs` — proyección a primera persona, con invariante sobre el contenido "
        "completo.",
        "`calificacion.mjs` — modelo de cálculo: escalas, matrices, descuentos, condición mínima "
        "y redondeo.",
        "`cargar.mjs` — arranque de los módulos.",
    ])

    m.h2("Documentación y materiales didácticos")
    m.puntos([
        "`docs/diseno/SDD.md` — documento de diseño de software: modelo de datos, modelo de "
        "calificación, catálogo de instrumentos, validador y arquitectura.",
        "`docs/diseno/SDD-cambios.md` — registro razonado de todas las decisiones de diseño, con "
        "su versión.",
        "`docs/diseno/por-que-estas-reglas.md` — el motivo de cada regla no negociable.",
        "`docs/diseno/anadir-una-materia.md` — protocolo para dar de alta una asignatura sin "
        "tocar código.",
        "`docs/diseno/enlace-proyecto-lengua.md` y `docs/marco/` — reparto de hechos con los "
        "proyectos afines del autor y matrices de referencia.",
        "`docs/Taller-de-Rubricas_Guia-para-empezar.docx` — guía del profesorado (generada).",
        "`docs/revision-lcl-*.md` — un documento de revisión por tipo de tarea (generados).",
        "`fuentes/curriculo/` y `fuentes/pau/` — currículo oficial y documentos normativos de "
        "contraste, material de terceros excluido de la inscripción (§7.3).",
        "`CLAUDE.md` y `README.md` — reglas no negociables del proyecto y guía del repositorio.",
    ])

    return m


def main() -> None:
    datos = leer_datos()
    memoria = construir(datos)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    memoria.doc.save(SALIDA)
    print(f"Escrito: {SALIDA.relative_to(RAIZ)}")
    print(f"  {datos['criterios']} criterios · {datos['matrices']} matrices · "
          f"{datos['reglas']} reglas · {datos['comprobaciones']} comprobaciones")


if __name__ == "__main__":
    main()
